import pandas as pd
import io
from flask import Blueprint, send_file, request, render_template
from flask_login import login_required
from models.db import db
from models.voluntarios.ponto import Ponto
from models.voluntarios.voluntario import Voluntario
from models.voluntarios.atividade import Atividade
from models.user.pessoa import Pessoa
from sqlalchemy import func, extract
from datetime import datetime
from openpyxl import Workbook
from io import BytesIO
import calendar

relatorios_bp = Blueprint("relatorios", __name__, template_folder="../views")

@relatorios_bp.route("/relatorios")
@login_required
def pagina_relatorios():
    hoje_str = datetime.utcnow().date().isoformat()
    return render_template("relatorios.html", hoje_str=hoje_str)

def _get_pivot_table_df(agg_column, agg_func):
    
    if agg_column == 'horas':
        query_data = db.session.query(
            Atividade.nome.label('Atividade'),
            Ponto.entrada,
            Ponto.saida,
            Ponto.duracao
        ).join(Atividade).filter(Ponto.saida != None).all()
        
        df_raw = pd.DataFrame(query_data)
        if df_raw.empty:
            return pd.DataFrame(columns=["Atividade"])
            
        df_raw['Horas'] = df_raw['duracao'].dt.total_seconds() / 3600
        df_raw['Mes'] = df_raw['entrada'].dt.month
        df_raw['Ano'] = df_raw['entrada'].dt.year
        
        pivot = pd.pivot_table(
            df_raw,
            values='Horas',
            index='Atividade',
            columns=['Ano', 'Mes'],
            aggfunc=agg_func,
            fill_value=0,
            margins=True,
            margins_name="Total Geral"
        )
    else:
        query_data = db.session.query(
            Atividade.nome.label('Atividade'),
            extract('month', Ponto.entrada).label('Mes'),
            extract('year', Ponto.entrada).label('Ano'),
            func.count(Ponto.id).label('Contagem')
        ).join(Atividade).filter(Ponto.saida != None)\
         .group_by('Atividade', 'Ano', 'Mes')\
         .order_by('Ano', 'Mes', 'Atividade').all()
        
        df_raw = pd.DataFrame(query_data)
        if df_raw.empty:
            return pd.DataFrame(columns=["Atividade"])
            
        pivot = pd.pivot_table(
            df_raw,
            values='Contagem',
            index='Atividade',
            columns=['Ano', 'Mes'],
            aggfunc=agg_func,
            fill_value=0,
            margins=True,
            margins_name="Total Geral"
        )
    
    return pivot

def _get_indicadores_voluntarios_df():
    entradas_query = db.session.query(
        extract('year', Voluntario.data_entrada).label('Ano'),
        extract('month', Voluntario.data_entrada).label('Mes'),
        func.count(Voluntario.id).label('Voluntários que Entraram')
    ).group_by('Ano', 'Mes').order_by('Ano', 'Mes')
    
    df_entradas = pd.DataFrame(entradas_query.all())

    atuantes_query = db.session.query(
        extract('year', Ponto.entrada).label('Ano'),
        extract('month', Ponto.entrada).label('Mes'),
        func.count(func.distinct(Ponto.id_voluntario)).label('Voluntários Atuantes')
    ).filter(Ponto.saida != None).group_by('Ano', 'Mes').order_by('Ano', 'Mes')
    
    df_atuantes = pd.DataFrame(atuantes_query.all())

    all_entries = db.session.query(Voluntario.data_entrada).all()
    if not all_entries:
        df_geral_final = pd.DataFrame(columns=['Ano', 'Mes'])
    else:
        df_all_entries = pd.DataFrame(all_entries, columns=['data_entrada'])
        df_all_entries['data_entrada'] = pd.to_datetime(df_all_entries['data_entrada'])
        df_all_entries = df_all_entries.set_index('data_entrada').assign(count=1)
        df_cumulative = df_all_entries.resample('ME').sum().cumsum()
        df_cumulative = df_cumulative.rename(columns={'count': 'Total de Voluntários (Acumulado)'})
        df_cumulative['Ano'] = df_cumulative.index.year
        df_cumulative['Mes'] = df_cumulative.index.month
        
        df_geral = pd.merge(df_entradas, df_atuantes, on=['Ano', 'Mes'], how='outer')
        df_geral_final = pd.merge(df_geral, df_cumulative.reset_index(drop=True), on=['Ano', 'Mes'], how='outer')
        df_geral_final = df_geral_final.fillna(0)
    
    try:
        df_geral_pivot = df_geral_final.set_index(['Ano', 'Mes']).T
    except KeyError:
        return pd.DataFrame(columns=["Indicador"])
    
    return df_geral_pivot

@relatorios_bp.route("/relatorio/lista_presenca_mes")
@login_required
def gerar_presenca_mensal():
    mes = int(request.args.get("mes", datetime.now().month))
    ano = datetime.now().year

    ultimo_dia = calendar.monthrange(ano, mes)[1]
    inicio = datetime(ano, mes, 1, 0, 0, 0)
    fim = datetime(ano, mes, ultimo_dia, 23, 59, 59)

    pontos = (
        Ponto.query
        .join(Voluntario)
        .join(Pessoa, Voluntario.pessoa)
        .join(Atividade)
        .filter(Ponto.entrada.between(inicio, fim))
        .all()
    )

    wb = Workbook()
    wb.remove(wb.active)

    atividades = {}
    for ponto in pontos:
        if ponto.atividade.nome not in atividades:
            atividades[ponto.atividade.nome] = []
        atividades[ponto.atividade.nome].append(ponto)

    for atividade_nome, pontos_atividade in atividades.items():
        ws = wb.create_sheet(title=atividade_nome[:31])
        ws.append(["Nome", "Atividade", "Data", "Entrada", "Saída", "Total de Hora"])

        for ponto in pontos_atividade:
            entrada = ponto.entrada.strftime("%d/%m/%Y %H:%M")
            saida = ponto.saida.strftime("%d/%m/%Y %H:%M") if ponto.saida else ""
            total_horas = round((ponto.saida - ponto.entrada).total_seconds() / 3600, 2) if ponto.saida else ""
            ws.append([
                ponto.voluntario.pessoa.nome,
                ponto.atividade.nome,
                ponto.entrada.strftime("%d/%m/%Y"),
                ponto.entrada.strftime("%H:%M"),
                ponto.saida.strftime("%H:%M") if ponto.saida else "",
                total_horas
            ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Presenca_{mes}_{ano}.xlsx"
    return send_file(
        output,
        download_name=filename,
        as_attachment=True,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@relatorios_bp.route("/relatorio/gerar_relatorio_dashboard_dia")
@login_required
def gerar_relatorio_dashboard_dia():
    date_str = request.args.get('date')
    if date_str:
        try:
            target_date = datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            target_date = datetime.utcnow().date()
    else:
        target_date = datetime.utcnow().date()

    pontos = Ponto.query.join(Voluntario).join(Voluntario.pessoa).join(Atividade)\
        .filter(Ponto.entrada >= datetime.combine(target_date, datetime.min.time()))\
        .filter(Ponto.entrada <= datetime.combine(target_date, datetime.max.time())).all()

    pontos_abertos = [p for p in pontos if not p.saida]

    kpis = {
        "Trabalhando Agora": len([p for p in pontos if not p.saida]),
        f"Entradas em {target_date}": len(pontos),
        f"Saídas em {target_date}": len([p for p in pontos if p.saida]),
        "Total Voluntários Ativos": Voluntario.query.filter_by(status="ativo").count(),
        "Total Atividades Ativas": Atividade.query.filter_by(ativo=True).count()
    }

    wb = Workbook()
    wb.remove(wb.active)

    ws_kpi = wb.create_sheet("Resumo_Dashboard_Dia")
    ws_kpi.append(["Métrica", "Valor"])
    for chave, valor in kpis.items():
        ws_kpi.append([chave, valor])

    for col_cells in ws_kpi.columns:
        max_length = max(len(str(cell.value)) for cell in col_cells)
        ws_kpi.column_dimensions[col_cells[0].column_letter].width = max_length + 2

    ws_abertos = wb.create_sheet("Trabalhando_Agora")
    ws_abertos.append(["Voluntário", "Entrada", "Origem"])
    for p in pontos_abertos:
        ws_abertos.append([
            p.voluntario.pessoa.nome,
            p.entrada.strftime('%Y-%m-%d %H:%M:%S'),
            p.origem or ""
        ])

    for col_cells in ws_abertos.columns:
        max_length = max(len(str(cell.value)) for cell in col_cells)
        ws_abertos.column_dimensions[col_cells[0].column_letter].width = max_length + 2

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=f'relatorio_pulse_dia_{target_date}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )